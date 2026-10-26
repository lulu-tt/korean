#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
지역어 기상도 원자료 폴더 검사 → 담당자 확인 요청 엑셀
=====================================================
새 자료를 받으면 올리기 전에 이 스크립트로 먼저 본다.

    python3 scripts/check_gisangdo_files.py "data/dialect_gisangdo_25(261026)"
    python3 scripts/check_gisangdo_files.py <폴더> -o data/exports/검사결과.xlsx

판정·대조 기준은 모두 저장소 안의 한 곳을 쓴다.
  · 서비스 목록  scripts/gisangdo_service_items.json   (101개, hidden 항목 포함)
  · 기준표      scripts/gisangdo_region_vocab.json    (지역별 방언형 → -B- 대조)
  · B 대조      etl_awareness_region.b_check()        (관리자 업로드 경고와 같은 함수)

줄 단위 결과는 원본과 같은 6칸(제보자 · 일련번호 · 항목번호 · 조사어휘 · 방언형(기저형) ·
사용도/인지도)에 파일명과 실제 엑셀 행 번호를 붙여 낸다 — 담당자가 원본에서 바로 찾게.
"""
import argparse, collections, datetime, os, re, sys, unicodedata

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import etl_awareness_region as E                                     # noqa: E402

N = lambda s: unicodedata.normalize('NFC', str(s)).strip() if s is not None else ''
FNAME = re.compile(r'^([A-Z]{2})(\d{2})(\d{2})([MF])VE')
COLS6 = ['제보자', '일련번호', '항목번호', '조사어휘', '방언형(기저형)', '사용도/인지도']
# 머리글 별칭 — neibis-cms/serve.py _wb_norm_header 와 같은 이름을 받는다
ALIAS = {'제보자': '제보자', '일련번호': '일련번호', '항목번호': '항목번호',
         '조사어휘': '조사어휘', '표제어': '조사어휘', '표제어형': '조사어휘', '표준어형': '조사어휘', '대응표준어': '조사어휘',
         '방언형(기저형)': '방언형(기저형)', '사용도/인지도': '사용도/인지도', '인지도/사용도': '사용도/인지도'}


def header_key(h):
    s = re.sub(r'\s+', '', N(h))
    s = re.sub(r'\((수정|보충|최종|검수)\)$', '', s)
    return ALIAS.get(s)


def read_folder(d):
    out = {}
    for f in sorted(os.listdir(d)):
        if not f.endswith('.xlsx') or f.startswith('~$'):
            continue
        name = N(f)
        m = FNAME.match(name)
        if not m:
            out[name] = {'file': name, 'bad_name': True}
            continue
        k = m.group(0)
        p = os.path.join(d, f)
        wb = openpyxl.load_workbook(p, data_only=True)
        wf = openpyxl.load_workbook(p)                       # 수식 칸 찾기용
        ws = wb.worksheets[0]
        hdr = [N(x) for x in next(ws.iter_rows(max_row=1, values_only=True), [])]
        idx = {}
        for i, h in enumerate(hdr):
            key = header_key(h)
            if key and key not in idx:
                idx[key] = i
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True), 1):
            if i == 1 or not r or all(v in (None, '') for v in r):
                continue
            rows.append((i, [N(r[idx[c]]) if c in idx and idx[c] < len(r) else '' for c in COLS6]))
        others = [s.title for s in wb.worksheets[1:]
                  if any(any(v not in (None, '') for v in row) for row in s.iter_rows(values_only=True))]
        nocache = sum(1 for row in wf.worksheets[0].iter_rows() for x in row
                      if isinstance(x.value, str) and x.value.startswith('=')
                      and ws[x.coordinate].value in (None, ''))
        out[k] = {'file': name, 'key': k, 'region': m.group(1), 'hdr': hdr, 'idx': idx,
                  'rows': rows, 'others': others, 'nocache': nocache}
    return out


def check(F):
    allow = E.load_service_items() or set()
    file_issues, row_issues = [], []
    for k, x in F.items():
        if x.get('bad_name'):
            file_issues.append((x['file'], '파일명이 규칙({지역2}{연차2}{세대2}{성별1}VE…)과 다름', '파일째 거부됩니다'))
            continue
        if not x['file'].endswith('_사용도인지도.xlsx'):
            file_issues.append((x['file'], '파일명 꼬리가 _사용도인지도 가 아님', '지장 없음'))
        miss = [c for c in COLS6 if c not in x['idx']]
        if miss:
            file_issues.append((x['file'], '첫 시트 머리글에 없는 칸: ' + ', '.join(miss)
                                + ' (머리글: ' + ' | '.join(h for h in x['hdr'] if h) + ')',
                                '조사 원본이거나 서식이 다릅니다' if set(miss) & {'항목번호', '조사어휘', '사용도/인지도'}
                                else '그 칸이 빈 값으로 들어갑니다'))
        if x['others']:
            file_issues.append((x['file'], '첫 시트 말고 값이 든 시트: ' + ', '.join(x['others']), '첫 시트만 읽습니다'))
        if x['nocache']:
            file_issues.append((x['file'], '결과값이 없는 수식 칸 %d개' % x['nocache'], '빈 값으로 읽힙니다'))
        who = {r[0] for _, r in x['rows'] if r[0]}
        if who and who != {k}:
            file_issues.append((x['file'], '제보자 칸이 파일명과 다름: ' + ', '.join(sorted(who - {k})[:3]), '파일째 거부됩니다'))
        same = collections.defaultdict(set)
        for i, r in x['rows']:
            sn, cd, hw, df, g = r[1], r[2], r[3], r[4], r[5]
            if not cd:
                row_issues.append(('항목번호 빈칸', k, i, r, '항목번호를 채우거나 줄을 빼 주십시오'))
                continue
            base = re.match(r'^(\d{5})', cd)
            if not base:
                row_issues.append(('항목번호 형식 이상', k, i, r, '5자리 항목번호(+하위코드)로 적어 주십시오'))
                continue
            if g not in ('1', '2', '3', '4'):
                row_issues.append(('사용도/인지도 이상', k, i, r, '1~4로 채우거나, 등급이 없으면 줄을 빼 주십시오'))
            if re.search(r'-[AB]-\d', hw):
                row_issues.append(('조사어휘 칸에 하위코드', k, i, r, '하위코드는 항목번호 칸에 적어 주십시오'))
            if allow and base.group(1) not in allow:
                row_issues.append(('서비스 목록 밖 항목번호', k, i, r, '번호 오타인지 확인해 주십시오'))
            if sn and sn != k + cd:
                row_issues.append(('일련번호와 항목번호가 다름', k, i, r, '일련번호(%s)와 항목번호(%s) 중 맞는 쪽으로' % (sn, cd)))
            if re.search(r'-[AB]-\d', cd):
                same[cd].add(hw)
        for cd, hws in same.items():
            if len(hws) > 1:
                for i, r in x['rows']:
                    if r[2] == cd:
                        row_issues.append(('같은 하위코드에 다른 말', k, i, r, '서로 다른 말은 다음 빈 번호로 나눠 주십시오'))
    # -B- 대조 (기준표)
    pairs = [(x['region'], r[2]) for x in F.values() if not x.get('bad_name') for _, r in x['rows']]
    bres = E.b_check(pairs)
    vocab = E.load_region_vocab()
    bdiff = [(k, i, r) for k, x in F.items() if not x.get('bad_name') for i, r in x['rows']
             if '-B-' in r[2] and r[4] not in ('', '*')
             and E.norm(r[4]) not in E.head_forms((vocab.get(r[2][:5]) or {}).get('word') or r[3])]
    return file_issues, row_issues, bres, bdiff


th = Side(style='thin', color='C9D1DC')
BOX = Border(left=th, right=th, top=th, bottom=th)
HEAD = PatternFill('solid', fgColor='1F3864')
SUG = PatternFill('solid', fgColor='EAF5EC')
EDIT = PatternFill('solid', fgColor='FFF7D6')
SOFT = PatternFill('solid', fgColor='EEF3FA')
RN = E.REGION_NAMES


def font(b=False, c='1F2937', s=10):
    return Font(name='맑은 고딕', bold=b, color=c, size=s)


def sheet(wb, title, sub, cols, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws['A1'] = title
    ws['A1'].font = font(True, '1F3864', 13)
    ws['A2'] = sub
    ws['A2'].font = font(c='4B5563')
    for i, (h, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(4, i, h)
        x.font = font(True, 'FFFFFF')
        x.fill = HEAD
        x.border = BOX
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[x.column_letter].width = w
    return ws


def put(ws, r, vals, fills=None, center=()):
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, N(v) if isinstance(v, str) else v)
        x.font = font()
        x.border = BOX
        x.alignment = Alignment(vertical='top', wrap_text=True, horizontal='center' if c in center else 'left')
        if fills and fills.get(c) is not None:
            x.fill = fills[c]


def write_xlsx(path, folder, F, file_issues, row_issues, bres, bdiff):
    wb = Workbook()
    s = wb.active
    s.title = '요약'
    s.sheet_view.showGridLines = False
    nrows = sum(len(x['rows']) for x in F.values() if not x.get('bad_name'))
    miss_regions = collections.Counter(x['kind'] for x in bres)
    lines = [
        ('지역어 기상도 원자료 검사 결과', True),
        ('대상: %s · 파일 %d개 · 줄 %d · 검사일 %s' % (folder, len(F), nrows, datetime.date.today()), False), ('', False),
        ('파일 단위 문제 %d건  → "파일"' % len(file_issues), True),
        ('줄 단위 문제 %d줄  → "줄"' % len(row_issues), True),
    ] + [('   · %s %d' % (k, v), False) for k, v in collections.Counter(t[0] for t in row_issues).most_common()] + [
        ('B 대조(기준표) — 기준표상 표준어와 다른 지역에 B 가 붙은 곳 %d곳  → "B 대조"' % miss_regions['B 지역 아님'], True),
        ('B 줄인데 방언형이 표준어와 다른 줄 %d  → "B 방언형 다름" (지역 단위로 붙인 것이면 그대로 둠)' % len(bdiff), False),
        ('', False), ('칸 색: 초록 = 제안 · 노랑 = 처리 결과를 적을 칸', False)]
    for i, (t, b) in enumerate(lines, 1):
        s.cell(i, 1, t).font = font(b, '1F3864' if b else '374151', 11 if b else 10)
    s.column_dimensions['A'].width = 120
    f = sheet(wb, '파일', '파일 단위 문제', ['번호', '파일', '문제', '그대로 올리면', '처리 결과'], [6, 34, 70, 26, 18])
    for n, (fn, p, eff) in enumerate(file_issues, 1):
        put(f, 4 + n, [n, fn, p, eff, ''], {5: EDIT}, center=(1,))
    q = sheet(wb, '줄', '줄 단위 — 파일 · 엑셀 행으로 찾으시면 됩니다',
              ['번호', '유형', '파일', '엑셀 행'] + COLS6 + ['제안', '처리 결과'], [6, 22, 32, 8, 12, 24, 12, 14, 16, 10, 40, 18])
    for n, (t, k, i, r, sg) in enumerate(sorted(row_issues, key=lambda t: (t[0], t[1], t[2])), 1):
        put(q, 4 + n, [n, t, F[k]['file'], i] + r + [sg, ''], {11: SUG, 12: EDIT}, center=(1, 4, 10))
    b = sheet(wb, 'B 대조', '기준표(지역별 방언형)와 -B- 표기 대조 — 기준표상 그 지역 방언형이 표준어와 다른데 B 가 붙은 곳. 지역 · 항목마다 그 지역 제보자 줄을 함께 실었습니다',
              ['번호', '구분', '지역', '항목번호', '조사어휘', '기준표의 지역 방언형', '파일', '엑셀 행'] + COLS6 + ['확인'],
              [6, 12, 7, 10, 12, 18, 32, 8, 12, 24, 12, 14, 16, 10, 16])
    r = 5
    for n, x in enumerate(bres, 1):
        rows = [(F[k]['file'], i, rr) for k in sorted(F) if not F[k].get('bad_name') and F[k]['region'] == x['rg']
                for i, rr in F[k]['rows'] if rr[2][:5] == x['it']]
        fill = SOFT if n % 2 == 0 else None
        for j, (fn, i, rr) in enumerate(rows):
            head = [n, x['kind'], RN.get(x['rg'], x['rg']), x['it'], x['word'], '/'.join(x['listed'])] if j == 0 else [''] * 6
            put(b, r, head + [fn, i] + rr + [''], {c: fill for c in range(1, 15)} | ({15: EDIT} if j == 0 else {}),
                center=(1, 3, 4, 8, 14))
            r += 1
    d = sheet(wb, 'B 방언형 다름', '-B- 줄인데 방언형이 표준어와 다른 줄', ['번호', '파일', '엑셀 행'] + COLS6 + ['확인'],
              [6, 32, 8, 12, 24, 12, 14, 16, 10, 16])
    for n, (k, i, rr) in enumerate(sorted(bdiff), 1):
        put(d, 4 + n, [n, F[k]['file'], i] + rr + [''], {10: EDIT}, center=(1, 3, 9))
    for ws in wb.worksheets[1:]:
        ws.freeze_panes = 'A5'
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description='지역어 기상도 원자료 폴더 검사')
    ap.add_argument('folder')
    ap.add_argument('-o', '--out', help='확인 요청 엑셀 경로 (기본: data/exports/기상도_검사_<폴더명>.xlsx)')
    a = ap.parse_args()
    F = read_folder(a.folder)
    file_issues, row_issues, bres, bdiff = check(F)
    print('파일 %d개 · 줄 %d' % (len(F), sum(len(x['rows']) for x in F.values() if not x.get('bad_name'))))
    print('파일 단위 %d · 줄 단위 %d · B 대조 %d · B 방언형 다름 %d' % (len(file_issues), len(row_issues), len(bres), len(bdiff)))
    for t, v in collections.Counter(t[0] for t in row_issues).most_common():
        print('  ·', t, v)
    for t, v in collections.Counter(x['kind'] for x in bres).most_common():
        print('  ·', t, v)
    out = a.out or os.path.join('data', 'exports', '기상도_검사_%s.xlsx' % N(os.path.basename(os.path.normpath(a.folder))))
    write_xlsx(out, a.folder, F, file_issues, row_issues, bres, bdiff)
    print(out)


if __name__ == '__main__':
    main()
